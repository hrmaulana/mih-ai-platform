from pathlib import Path

from .base import Segment


def _first_line(text: str) -> str | None:
    line = next((ln.strip() for ln in text.splitlines() if ln.strip()), None)
    return line or None


def _shape_texts(shape) -> list[str]:
    """Teks dari satu shape: text frame atau tabel (baris digabung dengan ' | ')."""
    if shape.has_text_frame:
        t = shape.text_frame.text.strip()
        return [t] if t else []
    if shape.has_table:
        rows: list[str] = []
        for row in shape.table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                rows.append(row_text)
        return rows
    return []


def _slide_text(slide) -> str:
    """Gabungan teks semua shape + catatan pembicara untuk satu slide."""
    texts: list[str] = []
    for shape in slide.shapes:
        texts.extend(_shape_texts(shape))
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            texts.append(f"[Catatan pembicara] {notes}")
    return "\n".join(texts).strip()


def parse_pptx(path: Path) -> list[Segment]:
    from pptx import Presentation

    prs = Presentation(str(path))
    segments: list[Segment] = []
    for i, slide in enumerate(prs.slides, start=1):
        text = _slide_text(slide)
        if text:
            segments.append(
                Segment(text=text, page_or_slide=i, section_title=_first_line(text))
            )
    return segments
